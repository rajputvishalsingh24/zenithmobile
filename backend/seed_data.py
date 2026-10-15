from pathlib import Path
import sys

# Support running both as `python -m backend.seed_data` and `python backend/seed_data.py`
BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from backend.database import get_db_connection, init_db, is_postgres

SEED_FILE_PATH = BASE_DIR / "data" / "events_seed.xlsx"

# Fallback seed data in case data/events_seed.xlsx is omitted in a serverless bundle
FALLBACK_SEED_EVENTS = [
    {
        "id": 1,
        "title": "HackZenith 24-Hour Hackathon",
        "category": "Technical",
        "description": "Build innovative AI, Web3, or sustainability solutions in a 24-hour coding sprint. Mentors from top tech companies will guide teams through ideation, prototyping, and live demos.",
        "event_date": "2026-11-14",
        "event_time": "09:00 AM",
        "venue": "Innovation Center - Hall A",
        "capacity": 120,
        "fee": 250,
        "icon": "💻"
    },
    {
        "id": 2,
        "title": "RoboWars: Arena Showdown",
        "category": "Technical",
        "description": "Custom-built combat robots go head-to-head in our armored polycarbonate arena. Compete in 15kg and 30kg weight classes for the ultimate engineering championship trophy.",
        "event_date": "2026-11-14",
        "event_time": "11:30 AM",
        "venue": "Central Quadrangle Arena",
        "capacity": 80,
        "fee": 300,
        "icon": "🤖"
    },
    {
        "id": 3,
        "title": "CodeGolf & Competitive Programming",
        "category": "Technical",
        "description": "Solve algorithmic brainteasers under intense time pressure using the shortest and fastest code possible. Open to Python, C++, and Java programmers.",
        "event_date": "2026-11-15",
        "event_time": "10:00 AM",
        "venue": "CS Block - Lab 301",
        "capacity": 60,
        "fee": 100,
        "icon": "⚡"
    },
    {
        "id": 4,
        "title": "Battle of the Bands: Decibel '26",
        "category": "Cultural",
        "description": "College rock, indie, and fusion bands compete live under festival stage lights with a concert-grade sound system. Original compositions earn bonus jury points.",
        "event_date": "2026-11-14",
        "event_time": "06:00 PM",
        "venue": "Open Air Amphitheatre",
        "capacity": 250,
        "fee": 200,
        "icon": "🎸"
    },
    {
        "id": 5,
        "title": "Nritya: Street & Contemporary Dance-Off",
        "category": "Cultural",
        "description": "High-energy solo and crew dance battles spanning hip-hop, contemporary, and Indo-fusion choreography. Judged by guest choreographers.",
        "event_date": "2026-11-14",
        "event_time": "03:00 PM",
        "venue": "Main Auditorium",
        "capacity": 180,
        "fee": 150,
        "icon": "💃"
    },
    {
        "id": 6,
        "title": "Stand-Up & Improv Comedy Night",
        "category": "Cultural",
        "description": "Five minutes on the mic to make the auditorium roar with laughter, followed by an interactive Whose-Line style improv showdown.",
        "event_date": "2026-11-15",
        "event_time": "05:00 PM",
        "venue": "Mini Auditorium",
        "capacity": 100,
        "fee": 0,
        "icon": "🎤"
    },
    {
        "id": 7,
        "title": "Futsal 5v5 Lightning Cup",
        "category": "Sports",
        "description": "Fast-paced 5-a-side turf football tournament with 15-minute halves, knockout brackets, and live commentary under the floodlights.",
        "event_date": "2026-11-14",
        "event_time": "04:00 PM",
        "venue": "North Turf Ground",
        "capacity": 80,
        "fee": 200,
        "icon": "⚽"
    },
    {
        "id": 8,
        "title": "3x3 Street Basketball Slam",
        "category": "Sports",
        "description": "Half-court 3x3 basketball tournament following official FIBA 3x3 rules, plus a crowd-favorite 3-point shootout contest.",
        "event_date": "2026-11-15",
        "event_time": "08:30 AM",
        "venue": "Outdoor Basketball Courts",
        "capacity": 64,
        "fee": 150,
        "icon": "🏀"
    },
    {
        "id": 9,
        "title": "Blitz Chess Championship",
        "category": "Sports",
        "description": "5+3 Swiss-system blitz chess tournament across 7 rounds. FIDE-rated arbiters and digital clocks provided for all boards.",
        "event_date": "2026-11-15",
        "event_time": "11:00 AM",
        "venue": "Student Activity Center - Room 102",
        "capacity": 50,
        "fee": 50,
        "icon": "♟️"
    },
    {
        "id": 10,
        "title": "Valorant LAN Champions Tour",
        "category": "Gaming",
        "description": "5v5 tactical shooter tournament on 240Hz LAN setups with zero-ping local servers and big-screen caster commentary for the finals.",
        "event_date": "2026-11-14",
        "event_time": "10:30 AM",
        "venue": "Esports Arena - Block B",
        "capacity": 100,
        "fee": 250,
        "icon": "🎮"
    },
    {
        "id": 11,
        "title": "EA FC 26 Console Showdown",
        "category": "Gaming",
        "description": "1v1 knockout tournament on PlayStation 5 consoles. Bring your best custom tactics and claim the Zenith Golden Controller.",
        "event_date": "2026-11-15",
        "event_time": "01:30 PM",
        "venue": "Esports Arena - Lounge",
        "capacity": 48,
        "fee": 100,
        "icon": "🕹️"
    },
    {
        "id": 12,
        "title": "Generative AI & Agentic Workflows Workshop",
        "category": "Workshop",
        "description": "Hands-on workshop building real LLM agents, RAG pipelines, and tool-calling applications from scratch. Bring your laptop!",
        "event_date": "2026-11-14",
        "event_time": "02:00 PM",
        "venue": "Seminar Hall 1",
        "capacity": 75,
        "fee": 0,
        "icon": "🧠"
    },
    {
        "id": 13,
        "title": "UI/UX & Product Design Sprint",
        "category": "Workshop",
        "description": "Learn rapid wireframing, Figma auto-layout mastery, and usability testing techniques used by top product teams.",
        "event_date": "2026-11-15",
        "event_time": "10:30 AM",
        "venue": "Design Studio - Floor 2",
        "capacity": 45,
        "fee": 0,
        "icon": "🎨"
    },
    {
        "id": 14,
        "title": "FPV Drone Building & Flight Clinic",
        "category": "Workshop",
        "description": "Understand quadcopter aerodynamics, flight controllers, and solder your first micro-drone frame followed by simulator flight training.",
        "event_date": "2026-11-15",
        "event_time": "02:30 PM",
        "venue": "Robotics Lab & South Lawn",
        "capacity": 35,
        "fee": 150,
        "icon": "🚁"
    }
]


def seed_events() -> int:
    """
    Seed the events table from data/events_seed.xlsx ("Events" sheet) or fallback seed list.
    Skips seeding if the events table already contains rows.
    """
    init_db()
    conn = get_db_connection()
    try:
        cursor = conn.cursor()
        cursor.execute("SELECT COUNT(*) FROM events;")
        existing_count = cursor.fetchone()[0]
        if existing_count > 0:
            return 0

        parsed_events = []

        if SEED_FILE_PATH.exists():
            try:
                import openpyxl
                wb = openpyxl.load_workbook(SEED_FILE_PATH, data_only=True)
                sheet = wb["Events"] if "Events" in wb.sheetnames else wb.active
                rows = list(sheet.iter_rows(values_only=True))

                for row in rows[1:]:
                    if not row or all(cell is None for cell in row):
                        continue
                    event_id, title, category, description, event_date, event_time, venue, capacity, fee, icon = (
                        row[0],
                        str(row[1]).strip() if row[1] is not None else "",
                        str(row[2]).strip() if row[2] is not None else "",
                        str(row[3]).strip() if row[3] is not None else "",
                        str(row[4]).strip() if row[4] is not None else "",
                        str(row[5]).strip() if row[5] is not None else "",
                        str(row[6]).strip() if row[6] is not None else "",
                        int(row[7]) if row[7] is not None else 0,
                        int(row[8]) if row[8] is not None else 0,
                        str(row[9]).strip() if len(row) > 9 and row[9] is not None else "✨",
                    )
                    if hasattr(row[4], "strftime"):
                        event_date = row[4].strftime("%Y-%m-%d")
                    if hasattr(row[5], "strftime"):
                        event_time = row[5].strftime("%H:%M")

                    parsed_events.append({
                        "id": int(event_id) if event_id is not None else None,
                        "title": title,
                        "category": category,
                        "description": description,
                        "event_date": event_date,
                        "event_time": event_time,
                        "venue": venue,
                        "capacity": capacity,
                        "fee": fee,
                        "icon": icon,
                    })
            except Exception:
                parsed_events = []

        if not parsed_events:
            parsed_events = FALLBACK_SEED_EVENTS

        inserted = 0
        for ev in parsed_events:
            if ev.get("id") is not None:
                cursor.execute(
                    """
                    INSERT INTO events (
                        id, title, category, description, event_date,
                        event_time, venue, capacity, fee, icon
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
                    """,
                    (
                        ev["id"],
                        ev["title"],
                        ev["category"],
                        ev["description"],
                        ev["event_date"],
                        ev["event_time"],
                        ev["venue"],
                        ev["capacity"],
                        ev["fee"],
                        ev["icon"],
                    ),
                )
            else:
                cursor.execute(
                    """
                    INSERT INTO events (
                        title, category, description, event_date,
                        event_time, venue, capacity, fee, icon
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?);
                    """,
                    (
                        ev["title"],
                        ev["category"],
                        ev["description"],
                        ev["event_date"],
                        ev["event_time"],
                        ev["venue"],
                        ev["capacity"],
                        ev["fee"],
                        ev["icon"],
                    ),
                )
            inserted += 1

        conn.commit()

        # In PostgreSQL, reset serial sequence so auto-increment proceeds from max(id)
        if is_postgres():
            try:
                cursor.execute(
                    "SELECT setval(pg_get_serial_sequence('events', 'id'), COALESCE(MAX(id), 1)) FROM events;"
                )
                conn.commit()
            except Exception:
                pass

        return inserted
    finally:
        conn.close()


def print_seeded_events() -> None:
    """Print all rows in the events table for verification."""
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    conn = get_db_connection()
    try:
        rows = conn.execute("SELECT * FROM events ORDER BY id ASC;").fetchall()
        print(f"\nTotal events in database: {len(rows)}")
        print("-" * 100)
        for r in rows:
            print(
                f"[{r['id']:02d}] {r['icon']} {r['title']} ({r['category']}) | "
                f"{r['event_date']} {r['event_time']} @ {r['venue']} | "
                f"Cap: {r['capacity']} | Fee: Rs.{r['fee']}"
            )
        print("-" * 100)
    finally:
        conn.close()


if __name__ == "__main__":
    count = seed_events()
    if count > 0:
        print(f"Successfully seeded {count} events.")
    else:
        print("Events table already contains rows; skipping seed.")
    print_seeded_events()
